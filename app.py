# -*- coding: utf-8 -*-
import os
import re
from dotenv import load_dotenv
import streamlit as st
import psycopg2
import pandas as pd
from io import BytesIO
from datetime import datetime

load_dotenv()

@st.cache_resource
def get_conn():
    return psycopg2.connect(
        host=os.getenv("DB_HOST"),
        port=os.getenv("DB_PORT", "5432"),
        database=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASS")
    )

def validar_mac(mac):
    if not mac:
        return True
    return bool(re.match(r'^([0-9A-Fa-f]{2}[:\-]){5}[0-9A-Fa-f]{2}$', mac.strip()))

def validar_email(email):
    if not email:
        return True
    return bool(re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email.strip()))

def validar_ip_format(ip):
    if not ip:
        return True
    parts = ip.strip().split('.')
    return len(parts) == 4 and all(p.isdigit() and 0 <= int(p) <= 255 for p in parts)

def safe_rollback():
    try:
        get_conn().rollback()
    except Exception:
        pass

def generar_excel_formateado(df, nombre_hoja="Datos"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    AZUL    = "1A3C5E"
    CELESTE = "EAF1FB"
    BORDE   = "BFBFBF"

    borde = Border(
        left=Side(style="thin", color=BORDE),
        right=Side(style="thin", color=BORDE),
        top=Side(style="thin", color=BORDE),
        bottom=Side(style="thin", color=BORDE),
    )

    headers = list(df.columns)
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font      = Font(bold=True, color="FFFFFF", size=10)
        c.fill      = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = borde
    ws.row_dimensions[1].height = 26

    fill_alt = PatternFill("solid", fgColor=CELESTE)
    for ri, (_, row) in enumerate(df.iterrows(), 2):
        for ci, h in enumerate(headers, 1):
            val = row[h]
            if val is None or (isinstance(val, float) and pd.isna(val)):
                val = ""
            c = ws.cell(row=ri, column=ci, value=val)
            c.font      = Font(size=9)
            c.alignment = Alignment(vertical="center")
            c.border    = borde
            if ri % 2 == 0:
                c.fill = fill_alt

    for ci, h in enumerate(headers, 1):
        vals = [str(h)] + [
            str(row[h]) if row[h] is not None and not (isinstance(row[h], float) and pd.isna(row[h])) else ""
            for _, row in df.iterrows()
        ]
        ws.column_dimensions[get_column_letter(ci)].width = min(max(len(v) for v in vals) + 4, 45)

    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# ════════════════════════════════════════
# FUNCIONES PDF
# ════════════════════════════════════════

def generar_pdf_generico(df, titulo, col_widths=None):
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4),
                            leftMargin=20, rightMargin=20,
                            topMargin=30, bottomMargin=20)
    styles = getSampleStyleSheet()
    elements = []

    titulo_par = Paragraph(
        f"<b>{titulo} — IMJ</b><br/>"
        f"<font size=9>Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}</font>",
        styles["Title"]
    )
    elements.append(titulo_par)
    elements.append(Spacer(1, 12))

    columnas = list(df.columns)
    data = [columnas]
    for _, row in df.iterrows():
        data.append([str(row.get(c, "") or "") for c in columnas])

    if col_widths is None:
        page_width = landscape(A4)[0] - 40
        col_widths = [page_width / len(columnas)] * len(columnas)

    tabla = Table(data, colWidths=col_widths, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, 0),  colors.HexColor("#1a3c5e")),
        ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
        ("FONTNAME",      (0, 0), (-1, 0),  "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, 0),  8),
        ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE",      (0, 1), (-1, -1), 7),
        ("ROWBACKGROUNDS",(0, 1), (-1, -1), [colors.white, colors.HexColor("#eaf1fb")]),
        ("GRID",          (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING",    (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(tabla)
    elements.append(Spacer(1, 10))
    elements.append(Paragraph(f"Total de registros: <b>{len(df)}</b>", styles["Normal"]))
    doc.build(elements)
    buffer.seek(0)
    return buffer


def generar_pdf_usuarios(df):
    col_widths = [30, 70, 90, 90, 80, 100, 130]
    return generar_pdf_generico(df, "Inventario de Usuarios", col_widths)


def generar_pdf_equipos_detalle(df, usuario_nombre=None):
    titulo = f"Equipos de Computo — {usuario_nombre}" if usuario_nombre else "Inventario de Equipos de Computo"
    col_widths = [80, 80, 100, 120, 160, 70]
    return generar_pdf_generico(df, titulo, col_widths)


def generar_pdf_equipos_resumen(df):
    col_widths = [140, 160, 180, 50]
    return generar_pdf_generico(df, "Resumen de Equipos de Computo", col_widths)


# ════════════════════════════════════════
# CONFIGURACIÓN STREAMLIT
# ════════════════════════════════════════

st.set_page_config(page_title="Sistema IMJ", layout="wide")

st.markdown("""
    <style>
    [data-testid="stElementToolbar"] { display: none !important; }
    </style>
""", unsafe_allow_html=True)

st.title("Sistema de Inventario IMJ")
st.markdown("---")

menu = st.sidebar.selectbox("Selecciona un modulo", [
    "🏠 Inicio",
    "👤 Usuarios",
    "💻 Equipos de Computo",
    "🌐 Direccionamiento IP",
    "📱 Telefonos",
    "🖨️ Impresoras",
    "📦 Insumos",
])

# ════════════════════════════════════════
# INICIO
# ════════════════════════════════════════
if menu == "🏠 Inicio":
    st.subheader("Bienvenido al Sistema de Inventario del IMJ")

    busqueda_global = st.text_input("🔍 Busqueda global", placeholder="Buscar usuario, IP, serie, extension, insumo...")

    if busqueda_global:
        try:
            conn = get_conn()
            cur = conn.cursor()
            term = f"%{busqueda_global}%"
            resultados = []

            cur.execute("""
                SELECT 'Usuario' as tipo, nombre || ' ' || apellido_paterno as descripcion,
                       COALESCE(puesto,'') as detalle, 'Usuarios' as modulo
                FROM usuarios WHERE activo = true
                AND (nombre ILIKE %s OR apellido_paterno ILIKE %s OR apellido_materno ILIKE %s OR correo ILIKE %s)
            """, (term, term, term, term))
            for r in cur.fetchall():
                resultados.append({"Tipo": r[0], "Descripcion": r[1], "Detalle": r[2], "Modulo": r[3]})

            cur.execute("""
                SELECT 'IP' as tipo, ip as descripcion,
                       COALESCE(usuario,'') || ' | ' || COALESCE(estatus,'') as detalle,
                       'Direccionamiento IP' as modulo
                FROM inventario_ips_completo
                WHERE ip ILIKE %s OR usuario ILIKE %s OR mac ILIKE %s
            """, (term, term, term))
            for r in cur.fetchall():
                resultados.append({"Tipo": r[0], "Descripcion": r[1], "Detalle": r[2], "Modulo": r[3]})

            cur.execute("""
                SELECT 'Equipo' as tipo, nombre_equipo as descripcion,
                       COALESCE(marca,'') || ' ' || COALESCE(modelo,'') || ' | Serie: ' || COALESCE(serie,'') as detalle,
                       'Equipos de Computo' as modulo
                FROM computo WHERE nombre_equipo ILIKE %s OR serie ILIKE %s OR mac_address ILIKE %s
            """, (term, term, term))
            for r in cur.fetchall():
                resultados.append({"Tipo": r[0], "Descripcion": r[1], "Detalle": r[2], "Modulo": r[3]})

            cur.execute("""
                SELECT 'Telefono' as tipo, extension::text as descripcion,
                       COALESCE(numero_general,'') as detalle, 'Telefonos' as modulo
                FROM telefonos WHERE extension::text ILIKE %s OR numero_general ILIKE %s
            """, (term, term))
            for r in cur.fetchall():
                resultados.append({"Tipo": r[0], "Descripcion": r[1], "Detalle": r[2], "Modulo": r[3]})

            cur.execute("""
                SELECT 'Impresora' as tipo, COALESCE(marca,'') || ' ' || COALESCE(modelo,'') as descripcion,
                       'Serie: ' || COALESCE(serie,'') || ' | IP: ' || COALESCE(ip_address::text,'') as detalle,
                       'Impresoras' as modulo
                FROM impresoras WHERE marca ILIKE %s OR modelo ILIKE %s OR serie ILIKE %s
            """, (term, term, term))
            for r in cur.fetchall():
                resultados.append({"Tipo": r[0], "Descripcion": r[1], "Detalle": r[2], "Modulo": r[3]})

            cur.execute("""
                SELECT 'Insumo' as tipo, nombre_insumo as descripcion,
                       'Stock: ' || stock_actual::text as detalle, 'Insumos' as modulo
                FROM insumos WHERE nombre_insumo ILIKE %s OR numero_parte ILIKE %s
            """, (term, term))
            for r in cur.fetchall():
                resultados.append({"Tipo": r[0], "Descripcion": r[1], "Detalle": r[2], "Modulo": r[3]})

            cur.close()

            if resultados:
                st.success(f"{len(resultados)} resultados para '{busqueda_global}'")
                st.dataframe(pd.DataFrame(resultados), use_container_width=True, hide_index=True)
            else:
                st.info(f"Sin resultados para '{busqueda_global}'.")
        except Exception as e:
            safe_rollback()
            st.error(f"Error en busqueda: {e}")
    else:
        try:
            conn = get_conn()
            cur  = conn.cursor()

            cur.execute("SELECT COUNT(*) FROM usuarios WHERE activo = true")
            total_usuarios = cur.fetchone()[0]
            cur.execute("SELECT COUNT(*) FROM computo")
            total_computo = cur.fetchone()[0]
            cur.execute("SELECT COUNT(*) FROM impresoras")
            total_impresoras = cur.fetchone()[0]
            cur.execute("SELECT COUNT(*) FROM computo WHERE estatus = 'activo'")
            equipos_activos = cur.fetchone()[0]
            cur.execute("SELECT COUNT(*) FROM computo WHERE estatus IN ('dañado','en reparacion')")
            equipos_danados = cur.fetchone()[0]
            cur.execute("SELECT COUNT(*) FROM inventario_ips_completo WHERE estatus ILIKE 'Libre%'")
            ips_libres = cur.fetchone()[0]
            cur.execute("SELECT COUNT(*) FROM inventario_ips_completo WHERE estatus ILIKE 'Ocupada%'")
            ips_ocupadas = cur.fetchone()[0]
            cur.execute("SELECT COUNT(*) FROM insumos WHERE stock_actual <= stock_minimo")
            insumos_bajos = cur.fetchone()[0]
            cur.execute("""
                SELECT d.nombre, COUNT(u.id_usuario)
                FROM departamentos d
                LEFT JOIN usuarios u ON u.id_departamento = d.id_departamento AND u.activo = true
                GROUP BY d.nombre ORDER BY COUNT(u.id_usuario) DESC
            """)
            df_deptos = pd.DataFrame(cur.fetchall(), columns=["Departamento", "Usuarios"])
            cur.close()

            st.markdown("#### Resumen general")
            col1, col2, col3, col4 = st.columns(4)
            col1.metric("👤 Usuarios activos",  total_usuarios)
            col2.metric("💻 Equipos",           total_computo)
            col3.metric("🖨️ Impresoras",        total_impresoras)
            col4.metric("📦 Insumos stock bajo", insumos_bajos,
                        delta=f"-{insumos_bajos}" if insumos_bajos > 0 else None,
                        delta_color="inverse")

            st.markdown("#### Equipos e IPs")
            col5, col6, col7, col8 = st.columns(4)
            col5.metric("✅ Equipos activos",     equipos_activos)
            col6.metric("⚠️ Dañados / Reparacion", equipos_danados,
                        delta=f"+{equipos_danados}" if equipos_danados > 0 else None,
                        delta_color="inverse")
            col7.metric("🟢 IPs libres",   ips_libres)
            col8.metric("🔴 IPs ocupadas", ips_ocupadas)

            st.markdown("#### Usuarios por departamento")
            st.bar_chart(df_deptos.set_index("Departamento"), use_container_width=True, height=280)

        except Exception as e:
            safe_rollback()
            st.error(f"Error al conectar con la BD: {e}")

# ════════════════════════════════════════
# USUARIOS
# ════════════════════════════════════════
elif menu == "👤 Usuarios":
    st.subheader("👤 Gestion de Usuarios")
    tab1, tab2, tab3, tab4 = st.tabs(["Gestionar", "Alta", "Baja", "Reactivar"])

    with tab1:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_departamento, nombre FROM departamentos ORDER BY nombre")
            deptos       = cur.fetchall()
            deptos_map   = {d[1]: d[0] for d in deptos}
            deptos_names = [d[1] for d in deptos]

            mostrar_inactivos = st.checkbox("Mostrar usuarios inactivos (dados de baja)")
            col_f1, col_f2, col_f3 = st.columns([1.5, 1, 2])
            depto_f    = col_f1.selectbox("Departamento", ["Todos"] + deptos_names, key="ver_u_depto")
            puesto_f   = col_f2.text_input("Puesto", placeholder="Ej: Jefe...", key="ver_u_puesto")
            busqueda_u = col_f3.text_input("Buscar", placeholder="Nombre, apellido, correo...", key="ver_u_busq")

            cur.execute("""
                SELECT u.id_usuario, u.nombre, u.apellido_paterno, u.apellido_materno,
                       u.puesto, u.correo, d.nombre as departamento
                FROM usuarios u
                LEFT JOIN departamentos d ON u.id_departamento = d.id_departamento
                WHERE u.activo = %s
                ORDER BY u.apellido_paterno
            """, (not mostrar_inactivos,))
            rows = cur.fetchall()
            cur.close()

            df = pd.DataFrame(rows, columns=[
                "id_usuario", "nombre", "ap_paterno", "ap_materno",
                "puesto", "correo", "departamento"
            ])

            if depto_f != "Todos":
                df = df[df["departamento"] == depto_f].reset_index(drop=True)
            if puesto_f:
                df = df[df["puesto"].str.contains(puesto_f, case=False, na=False)].reset_index(drop=True)
            if busqueda_u:
                mask = df[["nombre", "ap_paterno", "ap_materno", "correo"]].apply(
                    lambda col: col.astype(str).str.contains(busqueda_u, case=False, na=False)
                ).any(axis=1)
                df = df[mask].reset_index(drop=True)

            st.caption(f"{len(df)} usuarios — edita en la tabla y presiona **Guardar cambios**. Cambia Departamento para traspasar.")

            df_edited = st.data_editor(
                df,
                use_container_width=True,
                hide_index=True,
                num_rows="fixed",
                key="usuarios_data_editor",
                column_config={
                    "id_usuario":   None,
                    "nombre":       st.column_config.TextColumn("Nombre(s)"),
                    "ap_paterno":   st.column_config.TextColumn("Ap. Paterno"),
                    "ap_materno":   st.column_config.TextColumn("Ap. Materno"),
                    "puesto":       st.column_config.TextColumn("Puesto"),
                    "correo":       st.column_config.TextColumn("Correo"),
                    "departamento": st.column_config.SelectboxColumn("Departamento", options=deptos_names),
                },
            )

            col_btn, col_exp1, col_exp2 = st.columns([1, 1, 2])
            with col_btn:
                guardar = st.button("💾 Guardar cambios", type="primary", use_container_width=True, key="usr_guardar")
            with col_exp1:
                if not df_edited.empty:
                    _df_exp = df_edited.drop(columns=["id_usuario"], errors="ignore").rename(columns={
                        "nombre": "Nombre(s)", "ap_paterno": "Ap. Paterno", "ap_materno": "Ap. Materno",
                        "puesto": "Puesto", "correo": "Correo", "departamento": "Departamento",
                    })
                    excel_buf = generar_excel_formateado(_df_exp, "Usuarios")
                    st.download_button("📊 Excel", data=excel_buf.getvalue(),
                                       file_name=f"usuarios_IMJ_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       use_container_width=True, key="usr_export_excel")
            with col_exp2:
                if not df_edited.empty:
                    try:
                        df_pdf = df_edited.drop(columns=["id_usuario"], errors="ignore").rename(columns={
                            "nombre": "Nombre(s)", "ap_paterno": "Ap. Paterno",
                            "ap_materno": "Ap. Materno", "puesto": "Puesto",
                            "correo": "Correo", "departamento": "Departamento"
                        })
                        pdf_bytes = generar_pdf_usuarios(df_pdf).read()
                        st.download_button("📄 PDF", data=pdf_bytes,
                                           file_name=f"usuarios_IMJ_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
                                           mime="application/pdf", use_container_width=True, key="usr_export_pdf")
                    except Exception:
                        pass

            if guardar:
                errores = []
                for _, row in df_edited.iterrows():
                    correo_val = str(row.get("correo", "") or "").strip()
                    if correo_val and not validar_email(correo_val):
                        errores.append(f"Correo invalido: {correo_val}")
                if errores:
                    for err in errores:
                        st.warning(err)
                else:
                    try:
                        cur2 = get_conn().cursor()
                        count = 0
                        for _, row in df_edited.iterrows():
                            id_dep = deptos_map.get(row.get("departamento"))
                            cur2.execute("""
                                UPDATE usuarios
                                SET nombre=%s, apellido_paterno=%s, apellido_materno=%s,
                                    puesto=%s, correo=%s, id_departamento=%s
                                WHERE id_usuario=%s
                            """, (
                                row.get("nombre") or None,
                                row.get("ap_paterno") or None,
                                row.get("ap_materno") or None,
                                row.get("puesto") or None,
                                row.get("correo") or None,
                                id_dep,
                                int(row["id_usuario"]),
                            ))
                            count += 1
                        get_conn().commit()
                        cur2.close()
                        st.success(f"✅ {count} usuario(s) guardados correctamente.")
                        st.rerun()
                    except Exception as e:
                        safe_rollback()
                        st.error(f"Error al guardar: {e}")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab2:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_departamento, nombre FROM departamentos ORDER BY nombre")
            deptos = cur.fetchall()
            cur.close()
            with st.form("alta_usuario"):
                col1, col2, col3 = st.columns(3)
                nombre = col1.text_input("Nombre(s)")
                ap_pat = col2.text_input("Apellido Paterno")
                ap_mat = col3.text_input("Apellido Materno")
                col4, col5 = st.columns(2)
                puesto = col4.text_input("Puesto")
                correo = col5.text_input("Correo")
                depto_sel = st.selectbox("Departamento", [d[1] for d in deptos])
                id_depto  = next(d[0] for d in deptos if d[1] == depto_sel)
                if st.form_submit_button("✅ Dar de Alta"):
                    errores = []
                    if not nombre or not ap_pat:
                        errores.append("Llena al menos Nombre y Apellido Paterno.")
                    if correo and not validar_email(correo):
                        errores.append("El formato del correo no es valido.")
                    if errores:
                        for err in errores:
                            st.warning(err)
                    else:
                        cur2 = get_conn().cursor()
                        cur2.execute("""
                            INSERT INTO usuarios (nombre, apellido_paterno, apellido_materno, puesto, correo, id_departamento)
                            VALUES (%s, %s, %s, %s, %s, %s)
                        """, (nombre, ap_pat, ap_mat, puesto, correo, id_depto))
                        get_conn().commit()
                        cur2.close()
                        st.success(f"✅ Usuario {nombre} {ap_pat} dado de alta.")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab3:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_usuario, nombre, apellido_paterno FROM usuarios WHERE activo = true ORDER BY apellido_paterno")
            usuarios = cur.fetchall()
            cur.close()
            if not usuarios:
                st.info("No hay usuarios activos.")
            else:
                opciones = {f"{u[1]} {u[2]}": u[0] for u in usuarios}
                sel = st.selectbox("Selecciona el usuario a dar de baja", list(opciones.keys()), key="baja_u_sel")
                st.warning(f"⚠️ **{sel}** quedara inactivo. Sus datos se conservan y puede reactivarse.")
                confirmar = st.text_input("Escribe CONFIRMAR para continuar", key="baja_u_conf")
                if st.button("🔴 Dar de Baja", disabled=(confirmar.strip() != "CONFIRMAR")):
                    cur2 = get_conn().cursor()
                    cur2.execute("UPDATE usuarios SET activo = false WHERE id_usuario = %s", (opciones[sel],))
                    get_conn().commit()
                    cur2.close()
                    st.success(f"✅ {sel} dado de baja. Puedes reactivarlo desde la pestaña 'Reactivar'.")
                    st.rerun()
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab4:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_usuario, nombre, apellido_paterno FROM usuarios WHERE activo = false ORDER BY apellido_paterno")
            inactivos = cur.fetchall()
            cur.close()
            if not inactivos:
                st.info("No hay usuarios inactivos.")
            else:
                opciones = {f"{u[1]} {u[2]}": u[0] for u in inactivos}
                sel = st.selectbox("Selecciona el usuario a reactivar", list(opciones.keys()), key="react_u_sel")
                st.info(f"Se reactivara el acceso de **{sel}**.")
                if st.button("✅ Reactivar"):
                    cur2 = get_conn().cursor()
                    cur2.execute("UPDATE usuarios SET activo = true WHERE id_usuario = %s", (opciones[sel],))
                    get_conn().commit()
                    cur2.close()
                    st.success(f"✅ {sel} reactivado correctamente.")
                    st.rerun()
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

# ════════════════════════════════════════
# EQUIPOS DE COMPUTO
# ════════════════════════════════════════
elif menu == "💻 Equipos de Computo":
    st.subheader("💻 Equipos de Computo")
    tab1, tab2, tab3 = st.tabs(["Resumen", "Gestionar", "Agregar"])

    with tab1:
        try:
            conn = get_conn()
            cur = conn.cursor()
            if "expandido" not in st.session_state:
                st.session_state.expandido = False

            col_f1, col_f2, col_f3 = st.columns([1, 2, 0.4])
            filtro   = col_f1.selectbox("Estatus", ["Todos", "activo", "dañado", "en reparacion"], key="eq_filtro")
            busqueda = col_f2.text_input("Buscar usuario o equipo", placeholder="Ej: Adriana, Dell...", key="eq_busq")
            with col_f3:
                st.markdown("<br>", unsafe_allow_html=True)
                if st.button("⛶", help="Expandir tabla", key="eq_expand_top"):
                    st.session_state.expandido = not st.session_state.expandido
                    st.rerun()

            query_base = """
                SELECT
                    u.nombre || ' ' || u.apellido_paterno AS "Usuario",
                    STRING_AGG(DISTINCT NULLIF(c.nombre_equipo, ''), ', ') AS "Equipos",
                    STRING_AGG(DISTINCT NULLIF(c.serie, ''), ' / ') AS "Series",
                    COUNT(c.id_computo) AS "Total"
                FROM public.computo c
                INNER JOIN public.usuarios u ON c.id_usuario = u.id_usuario
            """
            if filtro == "Todos":
                cur.execute(query_base + " GROUP BY u.nombre, u.apellido_paterno ORDER BY \"Total\" DESC;")
            else:
                cur.execute(query_base + " WHERE c.estatus = %s GROUP BY u.nombre, u.apellido_paterno ORDER BY \"Total\" DESC;", (filtro,))

            df = pd.DataFrame(cur.fetchall(), columns=["Usuario", "Equipos", "Series", "Total"])
            cur.close()
            if busqueda:
                mask = df.apply(lambda col: col.astype(str).str.contains(busqueda, case=False, na=False)).any(axis=1)
                df = df[mask].reset_index(drop=True)

            altura = 700 if st.session_state.expandido else 420
            seleccion = st.dataframe(df, use_container_width=True, hide_index=True,
                                     height=altura, on_select="rerun", selection_mode="single-row")

            df_detalles = None
            usuario_sel = None

            if seleccion.selection.rows:
                indice = seleccion.selection.rows[0]
                usuario_sel = df.iloc[indice]["Usuario"]
                st.markdown("---")
                st.subheader(f"📋 Equipos de: {usuario_sel}")
                cur2 = get_conn().cursor()
                cur2.execute("""
                    SELECT c.nombre_equipo AS "Equipo", c.marca AS "Marca", c.modelo AS "Modelo",
                           c.serie AS "Serie", c.mac_address AS "MAC", c.estatus AS "Estatus"
                    FROM public.computo c
                    JOIN public.usuarios u ON c.id_usuario = u.id_usuario
                    WHERE (u.nombre || ' ' || u.apellido_paterno) = %s
                """, (usuario_sel,))
                df_detalles = pd.DataFrame(cur2.fetchall(), columns=["Equipo", "Marca", "Modelo", "Serie", "MAC", "Estatus"])
                cur2.close()
                st.table(df_detalles)

            tb1, tb2, tb3_col, tb4, _ = st.columns([0.8, 0.8, 0.5, 0.5, 6])
            with tb1:
                try:
                    if df_detalles is not None and not df_detalles.empty:
                        pdf_bytes = generar_pdf_equipos_detalle(df_detalles, usuario_sel).read()
                        nombre_archivo = f"equipos_{usuario_sel.replace(' ','_')}_{datetime.now().strftime('%Y%m%d')}.pdf"
                    else:
                        pdf_bytes = generar_pdf_equipos_resumen(df).read()
                        nombre_archivo = f"equipos_resumen_{datetime.now().strftime('%Y%m%d')}.pdf"
                    st.download_button("📄 PDF", data=pdf_bytes, file_name=nombre_archivo, mime="application/pdf", use_container_width=True)
                except Exception as e:
                    st.button("📄 PDF", disabled=True)
                    st.caption(f"Error PDF: {e}")
            with tb2:
                try:
                    if df_detalles is not None and not df_detalles.empty:
                        excel_buffer = generar_excel_formateado(df_detalles, "Equipos")
                        nombre_archivo = f"equipos_{usuario_sel.replace(' ','_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
                    else:
                        excel_buffer = generar_excel_formateado(df, "Resumen")
                        nombre_archivo = f"equipos_resumen_{datetime.now().strftime('%Y%m%d')}.xlsx"
                    st.download_button("📊 Excel", data=excel_buffer.getvalue(), file_name=nombre_archivo,
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       use_container_width=True)
                except Exception as e:
                    st.button("📊 Excel", disabled=True)
            with tb4:
                if st.button("⛶", help="Pantalla completa", key="eq_expand_bot"):
                    st.session_state.expandido = not st.session_state.expandido
                    st.rerun()

            if not seleccion.selection.rows:
                st.caption(f"💡 Selecciona una fila para ver detalles. Total: {len(df)} usuarios.")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab2:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_usuario, nombre || ' ' || apellido_paterno FROM usuarios ORDER BY apellido_paterno")
            rows_u = cur.fetchall()

            col_f1, col_f2 = st.columns([1, 2])
            estatus_f = col_f1.selectbox("Estatus", ["Todos", "activo", "dañado", "en reparacion"], key="gest_eq_est")
            busqueda_eq = col_f2.text_input("Buscar equipo o usuario", placeholder="Ej: Dell, Juan...", key="gest_eq_busq")

            query = """
                SELECT c.id_computo,
                       COALESCE(u.nombre || ' ' || u.apellido_paterno, '') as usuario,
                       c.nombre_equipo, c.marca, c.modelo, c.serie, c.mac_address, c.estatus
                FROM computo c LEFT JOIN usuarios u ON c.id_usuario = u.id_usuario
                WHERE 1=1
            """
            params = []
            if estatus_f != "Todos":
                query += " AND c.estatus = %s"
                params.append(estatus_f)
            query += " ORDER BY u.apellido_paterno, c.nombre_equipo"
            cur.execute(query, params)
            rows = cur.fetchall()
            cur.close()

            usuarios_map   = {r[1]: r[0] for r in rows_u}
            usuarios_names = list(usuarios_map.keys())

            df = pd.DataFrame(rows, columns=[
                "id_computo", "usuario", "nombre_equipo", "marca", "modelo", "serie", "mac", "estatus"
            ])

            if busqueda_eq:
                mask = df.apply(lambda col: col.astype(str).str.contains(busqueda_eq, case=False, na=False)).any(axis=1)
                df = df[mask].reset_index(drop=True)

            st.caption(f"{len(df)} equipos — edita en la tabla y presiona **Guardar cambios**. Cambia Usuario para reasignar.")

            df_edited = st.data_editor(
                df,
                use_container_width=True,
                hide_index=True,
                num_rows="fixed",
                key="eq_data_editor",
                column_config={
                    "id_computo":    None,
                    "usuario":       st.column_config.SelectboxColumn("Usuario", options=usuarios_names, required=True),
                    "nombre_equipo": st.column_config.TextColumn("Equipo"),
                    "marca":         st.column_config.TextColumn("Marca"),
                    "modelo":        st.column_config.TextColumn("Modelo"),
                    "serie":         st.column_config.TextColumn("Serie"),
                    "mac":           st.column_config.TextColumn("MAC"),
                    "estatus":       st.column_config.SelectboxColumn("Estatus",
                                        options=["activo", "dañado", "en reparacion"], required=True),
                },
            )

            col_btn, col_exp = st.columns([1, 4])
            with col_btn:
                guardar_eq = st.button("💾 Guardar cambios", type="primary", use_container_width=True, key="eq_guardar")
            with col_exp:
                if not df_edited.empty:
                    _df_exp = df_edited.drop(columns=["id_computo"], errors="ignore").rename(columns={
                        "usuario": "Usuario", "nombre_equipo": "Equipo", "marca": "Marca",
                        "modelo": "Modelo", "serie": "Serie", "mac": "MAC Address", "estatus": "Estatus",
                    })
                    excel_buf = generar_excel_formateado(_df_exp, "Equipos")
                    st.download_button("📥 Exportar Excel", data=excel_buf.getvalue(),
                                       file_name=f"equipos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       use_container_width=True, key="eq_export")

            if guardar_eq:
                errores = []
                for _, row in df_edited.iterrows():
                    mac_val = str(row.get("mac", "") or "").strip()
                    if mac_val and not validar_mac(mac_val):
                        errores.append(f"MAC invalido: {mac_val}")
                if errores:
                    for err in errores:
                        st.warning(err)
                else:
                    try:
                        cur3 = get_conn().cursor()
                        count = 0
                        for _, row in df_edited.iterrows():
                            id_usuario = usuarios_map.get(row.get("usuario"))
                            cur3.execute("""
                                UPDATE computo
                                SET id_usuario=%s, nombre_equipo=%s, marca=%s,
                                    modelo=%s, serie=%s, mac_address=%s, estatus=%s
                                WHERE id_computo=%s
                            """, (
                                id_usuario,
                                row.get("nombre_equipo") or None,
                                row.get("marca") or None,
                                row.get("modelo") or None,
                                row.get("serie") or None,
                                row.get("mac") or None,
                                row.get("estatus"),
                                int(row["id_computo"]),
                            ))
                            count += 1
                        get_conn().commit()
                        cur3.close()
                        st.success(f"✅ {count} equipo(s) guardados correctamente.")
                        st.rerun()
                    except Exception as e:
                        safe_rollback()
                        st.error(f"Error al guardar: {e}")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab3:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_usuario, nombre, apellido_paterno FROM usuarios WHERE activo = true ORDER BY apellido_paterno")
            usuarios = cur.fetchall()
            cur.close()
            opciones_usr = {f"{u[1]} {u[2]}": u[0] for u in usuarios}
            with st.form("agregar_equipo"):
                sel_usr    = st.selectbox("Asignar a usuario", list(opciones_usr.keys()))
                col1, col2 = st.columns(2)
                nombre_eq  = col1.text_input("Nombre del equipo")
                marca      = col2.text_input("Marca")
                col3, col4 = st.columns(2)
                modelo = col3.text_input("Modelo")
                serie  = col4.text_input("Serie")
                col5, col6 = st.columns(2)
                mac     = col5.text_input("MAC Address")
                estatus = col6.selectbox("Estatus", ["activo", "dañado", "en reparacion"])
                if st.form_submit_button("✅ Agregar equipo"):
                    errores = []
                    if not nombre_eq or not marca:
                        errores.append("Completa al menos Nombre del equipo y Marca.")
                    if mac and not validar_mac(mac):
                        errores.append("Formato de MAC no valido. Usa XX:XX:XX:XX:XX:XX")
                    if errores:
                        for err in errores:
                            st.warning(err)
                    else:
                        cur2 = get_conn().cursor()
                        cur2.execute("""
                            INSERT INTO computo (id_usuario, nombre_equipo, marca, modelo, serie, mac_address, estatus)
                            VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """, (opciones_usr[sel_usr], nombre_eq, marca, modelo, serie, mac, estatus))
                        get_conn().commit()
                        cur2.close()
                        st.success(f"✅ Equipo **{nombre_eq}** registrado y asignado a {sel_usr}.")
                        st.rerun()
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

# ════════════════════════════════════════
# TELEFONOS
# ════════════════════════════════════════
elif menu == "📱 Telefonos":
    st.subheader("📱 Telefonos")
    tab1, tab2 = st.tabs(["Gestionar", "Agregar"])

    with tab1:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_usuario, nombre || ' ' || apellido_paterno FROM usuarios ORDER BY apellido_paterno")
            rows_u = cur.fetchall()
            usuarios_map   = {r[1]: r[0] for r in rows_u}
            usuarios_names = list(usuarios_map.keys())

            col_f1, col_f2 = st.columns([1, 2])
            ext_f    = col_f1.text_input("Buscar extension", placeholder="Ej: 1234...", key="tel_ext_f")
            nombre_f = col_f2.text_input("Buscar usuario",   placeholder="Ej: Juan...",  key="tel_nom_f")

            cur.execute("""
                SELECT t.id_telefono, t.numero_general, t.extension::text,
                       COALESCE(u.nombre || ' ' || u.apellido_paterno, '') as usuario
                FROM telefonos t LEFT JOIN usuarios u ON t.id_usuario = u.id_usuario
                ORDER BY u.apellido_paterno
            """)
            rows = cur.fetchall()
            cur.close()

            df = pd.DataFrame(rows, columns=["id_telefono", "numero_general", "extension", "usuario"])

            if ext_f:
                df = df[df["extension"].astype(str).str.contains(ext_f, case=False, na=False)].reset_index(drop=True)
            if nombre_f:
                df = df[df["usuario"].str.contains(nombre_f, case=False, na=False)].reset_index(drop=True)

            st.caption(f"{len(df)} telefonos — edita en la tabla y presiona **Guardar cambios**. Cambia Usuario para reasignar.")

            df_edited = st.data_editor(
                df,
                use_container_width=True,
                hide_index=True,
                num_rows="fixed",
                key="tel_data_editor",
                column_config={
                    "id_telefono":    None,
                    "numero_general": st.column_config.TextColumn("Numero General"),
                    "extension":      st.column_config.TextColumn("Extension"),
                    "usuario":        st.column_config.SelectboxColumn("Usuario", options=usuarios_names),
                },
            )

            col_btn, col_exp = st.columns([1, 4])
            with col_btn:
                guardar_tel = st.button("💾 Guardar cambios", type="primary", use_container_width=True, key="tel_guardar")
            with col_exp:
                if not df_edited.empty:
                    _df_exp = df_edited.drop(columns=["id_telefono"], errors="ignore").rename(columns={
                        "numero_general": "Numero General", "extension": "Extension", "usuario": "Usuario",
                    })
                    excel_buf = generar_excel_formateado(_df_exp, "Telefonos")
                    st.download_button("📥 Exportar Excel", data=excel_buf.getvalue(),
                                       file_name=f"telefonos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       use_container_width=True, key="tel_export")

            if guardar_tel:
                try:
                    cur2 = get_conn().cursor()
                    count = 0
                    for _, row in df_edited.iterrows():
                        id_usuario = usuarios_map.get(row.get("usuario"))
                        cur2.execute("""
                            UPDATE telefonos
                            SET numero_general=%s, extension=%s, id_usuario=%s
                            WHERE id_telefono=%s
                        """, (
                            row.get("numero_general") or None,
                            row.get("extension") or None,
                            id_usuario,
                            int(row["id_telefono"]),
                        ))
                        count += 1
                    get_conn().commit()
                    cur2.close()
                    st.success(f"✅ {count} telefono(s) guardados correctamente.")
                    st.rerun()
                except Exception as e:
                    safe_rollback()
                    st.error(f"Error al guardar: {e}")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab2:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_usuario, nombre, apellido_paterno FROM usuarios ORDER BY apellido_paterno")
            usuarios = cur.fetchall()
            cur.close()
            opciones_usr = {f"{u[1]} {u[2]}": u[0] for u in usuarios}
            with st.form("agregar_telefono"):
                numero    = st.text_input("Numero general", value="(55)1500 1300")
                extension = st.text_input("Extension")
                sel_usr   = st.selectbox("Asignar a usuario", list(opciones_usr.keys()))
                if st.form_submit_button("✅ Agregar"):
                    if extension:
                        cur2 = get_conn().cursor()
                        cur2.execute("INSERT INTO telefonos (numero_general, extension, id_usuario) VALUES (%s, %s, %s)",
                                     (numero, extension, opciones_usr[sel_usr]))
                        get_conn().commit()
                        cur2.close()
                        st.success("✅ Telefono agregado correctamente.")
                    else:
                        st.warning("Por favor escribe la extension.")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

# ════════════════════════════════════════
# IMPRESORAS
# ════════════════════════════════════════
elif menu == "🖨️ Impresoras":
    st.subheader("🖨️ Impresoras")
    tab1, tab2 = st.tabs(["Gestionar", "Agregar"])

    with tab1:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_usuario, nombre || ' ' || apellido_paterno FROM usuarios ORDER BY apellido_paterno")
            rows_u = cur.fetchall()
            usuarios_map   = {r[1]: r[0] for r in rows_u}
            usuarios_names = list(usuarios_map.keys())

            cur.execute("""
                SELECT i.id_impresora, i.marca, i.modelo, i.serie,
                       COALESCE(i.ip_address::text, '') as ip,
                       COALESCE(i.firmware, '') as firmware,
                       COALESCE(u.nombre || ' ' || u.apellido_paterno, '') as usuario
                FROM impresoras i LEFT JOIN usuarios u ON i.id_usuario = u.id_usuario
                ORDER BY u.apellido_paterno, i.marca
            """)
            rows = cur.fetchall()
            cur.close()

            df = pd.DataFrame(rows, columns=["id_impresora", "marca", "modelo", "serie", "ip", "firmware", "usuario"])

            marcas_uniq = ["Todas"] + sorted(df["marca"].dropna().unique().tolist())
            col_f1, col_f2, col_f3 = st.columns([1, 1, 2])
            marca_f   = col_f1.selectbox("Marca", marcas_uniq, key="imp_marca_f")
            ip_f      = col_f2.text_input("IP", placeholder="Ej: 192.168...", key="imp_ip_f")
            usuario_f = col_f3.text_input("Buscar usuario", placeholder="Nombre o apellido...", key="imp_usr_f")

            if marca_f != "Todas":
                df = df[df["marca"] == marca_f].reset_index(drop=True)
            if ip_f:
                df = df[df["ip"].str.contains(ip_f, case=False, na=False)].reset_index(drop=True)
            if usuario_f:
                df = df[df["usuario"].str.contains(usuario_f, case=False, na=False)].reset_index(drop=True)

            st.caption(f"{len(df)} impresoras — edita en la tabla y presiona **Guardar cambios**. Cambia Usuario para reasignar.")

            df_edited = st.data_editor(
                df,
                use_container_width=True,
                hide_index=True,
                num_rows="fixed",
                key="imp_data_editor",
                column_config={
                    "id_impresora": None,
                    "marca":        st.column_config.TextColumn("Marca"),
                    "modelo":       st.column_config.TextColumn("Modelo"),
                    "serie":        st.column_config.TextColumn("Serie"),
                    "ip":           st.column_config.TextColumn("IP"),
                    "firmware":     st.column_config.TextColumn("Firmware"),
                    "usuario":      st.column_config.SelectboxColumn("Usuario", options=usuarios_names),
                },
            )

            col_btn, col_exp = st.columns([1, 4])
            with col_btn:
                guardar_imp = st.button("💾 Guardar cambios", type="primary", use_container_width=True, key="imp_guardar")
            with col_exp:
                if not df_edited.empty:
                    _df_exp = df_edited.drop(columns=["id_impresora"], errors="ignore").rename(columns={
                        "marca": "Marca", "modelo": "Modelo", "serie": "Serie",
                        "ip": "IP Address", "firmware": "Firmware", "usuario": "Usuario",
                    })
                    excel_buf = generar_excel_formateado(_df_exp, "Impresoras")
                    st.download_button("📥 Exportar Excel", data=excel_buf.getvalue(),
                                       file_name=f"impresoras_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       use_container_width=True, key="imp_export")

            if guardar_imp:
                errores = []
                for _, row in df_edited.iterrows():
                    ip_val = str(row.get("ip", "") or "").strip()
                    if ip_val and not validar_ip_format(ip_val):
                        errores.append(f"IP invalida: {ip_val}")
                if errores:
                    for err in errores:
                        st.warning(err)
                else:
                    try:
                        cur2 = get_conn().cursor()
                        count = 0
                        for _, row in df_edited.iterrows():
                            id_usuario = usuarios_map.get(row.get("usuario"))
                            ip_db = str(row.get("ip") or "").strip() or None
                            cur2.execute("""
                                UPDATE impresoras
                                SET marca=%s, modelo=%s, serie=%s,
                                    ip_address=%s, firmware=%s, id_usuario=%s
                                WHERE id_impresora=%s
                            """, (
                                row.get("marca") or None,
                                row.get("modelo") or None,
                                row.get("serie") or None,
                                ip_db,
                                row.get("firmware") or None,
                                id_usuario,
                                int(row["id_impresora"]),
                            ))
                            count += 1
                        get_conn().commit()
                        cur2.close()
                        st.success(f"✅ {count} impresora(s) guardadas correctamente.")
                        st.rerun()
                    except Exception as e:
                        safe_rollback()
                        st.error(f"Error al guardar: {e}")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab2:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_usuario, nombre, apellido_paterno FROM usuarios ORDER BY apellido_paterno")
            usuarios = cur.fetchall()
            cur.close()
            opciones_usr = {f"{u[1]} {u[2]}": u[0] for u in usuarios}
            with st.form("agregar_impresora"):
                col1, col2 = st.columns(2)
                marca   = col1.text_input("Marca")
                modelo  = col2.text_input("Modelo")
                col3, col4 = st.columns(2)
                serie   = col3.text_input("Serie")
                ip      = col4.text_input("IP Address")
                firmware = st.text_input("Firmware")
                sel_usr  = st.selectbox("Asignar a usuario", list(opciones_usr.keys()))
                if st.form_submit_button("✅ Agregar"):
                    errores = []
                    if not marca or not modelo:
                        errores.append("Llena al menos Marca y Modelo.")
                    if ip and not validar_ip_format(ip):
                        errores.append("Formato de IP no valido. Usa X.X.X.X")
                    if errores:
                        for err in errores:
                            st.warning(err)
                    else:
                        cur2 = get_conn().cursor()
                        cur2.execute("""
                            INSERT INTO impresoras (marca, modelo, serie, ip_address, firmware, id_usuario)
                            VALUES (%s, %s, %s, %s, %s, %s)
                        """, (marca, modelo, serie, ip, firmware, opciones_usr[sel_usr]))
                        get_conn().commit()
                        cur2.close()
                        st.success(f"✅ Impresora {marca} {modelo} agregada.")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

# ════════════════════════════════════════
# INSUMOS
# ════════════════════════════════════════
elif menu == "📦 Insumos":
    st.subheader("📦 Insumos")
    tab1, tab2, tab3 = st.tabs(["Gestionar", "Agregar", "Suministro"])

    with tab1:
        try:
            cur = get_conn().cursor()
            cur.execute("""
                SELECT id_insumo, nombre_insumo, numero_parte,
                       stock_minimo, stock_maximo, stock_actual
                FROM insumos ORDER BY nombre_insumo
            """)
            rows = cur.fetchall()
            cur.close()

            df = pd.DataFrame(rows, columns=[
                "id_insumo", "nombre", "numero_parte", "stock_min", "stock_max", "stock_actual"
            ])

            col_f1, col_f2 = st.columns([2, 1])
            busqueda_ins = col_f1.text_input("Buscar insumo", placeholder="Nombre o numero de parte...", key="ins_busq")
            stock_f = col_f2.selectbox("Stock", ["Todos", "Stock bajo / agotado", "Stock normal"], key="ins_stock_f")

            if busqueda_ins:
                mask = df[["nombre", "numero_parte"]].apply(
                    lambda col: col.astype(str).str.contains(busqueda_ins, case=False, na=False)
                ).any(axis=1)
                df = df[mask].reset_index(drop=True)
            if stock_f == "Stock bajo / agotado":
                df = df[df["stock_actual"] <= df["stock_min"]].reset_index(drop=True)
            elif stock_f == "Stock normal":
                df = df[df["stock_actual"] > df["stock_min"]].reset_index(drop=True)

            st.caption(f"{len(df)} insumos — edita nombre, numero de parte y niveles min/max. El stock actual se gestiona en la pestaña Suministro.")

            df_edited = st.data_editor(
                df,
                use_container_width=True,
                hide_index=True,
                num_rows="fixed",
                key="ins_data_editor",
                column_config={
                    "id_insumo":    None,
                    "nombre":       st.column_config.TextColumn("Insumo"),
                    "numero_parte": st.column_config.TextColumn("No. Parte"),
                    "stock_min":    st.column_config.NumberColumn("Stock Min",  min_value=0, step=1),
                    "stock_max":    st.column_config.NumberColumn("Stock Max",  min_value=0, step=1),
                    "stock_actual": st.column_config.NumberColumn("Stock Actual", disabled=True),
                },
            )

            col_btn, col_exp = st.columns([1, 4])
            with col_btn:
                guardar_ins = st.button("💾 Guardar cambios", type="primary", use_container_width=True, key="ins_guardar")
            with col_exp:
                if not df_edited.empty:
                    _df_exp = df_edited.drop(columns=["id_insumo"], errors="ignore").rename(columns={
                        "nombre": "Insumo", "numero_parte": "No. Parte",
                        "stock_min": "Stock Min", "stock_max": "Stock Max", "stock_actual": "Stock Actual",
                    })
                    excel_buf = generar_excel_formateado(_df_exp, "Insumos")
                    st.download_button("📥 Exportar Excel", data=excel_buf.getvalue(),
                                       file_name=f"insumos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       use_container_width=True, key="ins_export")

            if guardar_ins:
                try:
                    cur2 = get_conn().cursor()
                    count = 0
                    for _, row in df_edited.iterrows():
                        cur2.execute("""
                            UPDATE insumos
                            SET nombre_insumo=%s, numero_parte=%s,
                                stock_minimo=%s, stock_maximo=%s
                            WHERE id_insumo=%s
                        """, (
                            row.get("nombre") or None,
                            row.get("numero_parte") or None,
                            int(row.get("stock_min") or 0),
                            int(row.get("stock_max") or 0),
                            int(row["id_insumo"]),
                        ))
                        count += 1
                    get_conn().commit()
                    cur2.close()
                    st.success(f"✅ {count} insumo(s) guardados correctamente.")
                    st.rerun()
                except Exception as e:
                    safe_rollback()
                    st.error(f"Error al guardar: {e}")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab2:
        try:
            with st.form("agregar_insumo"):
                nombre    = st.text_input("Nombre del insumo")
                num_parte = st.text_input("Numero de parte")
                col1, col2, col3 = st.columns(3)
                stock_min = col1.number_input("Stock minimo", min_value=0, value=5)
                stock_max = col2.number_input("Stock maximo", min_value=0, value=50)
                stock_act = col3.number_input("Stock actual",  min_value=0, value=0)
                if st.form_submit_button("✅ Agregar"):
                    if nombre:
                        cur2 = get_conn().cursor()
                        cur2.execute("""
                            INSERT INTO insumos (nombre_insumo, numero_parte, stock_minimo, stock_maximo, stock_actual)
                            VALUES (%s, %s, %s, %s, %s)
                        """, (nombre, num_parte, stock_min, stock_max, stock_act))
                        get_conn().commit()
                        cur2.close()
                        st.success(f"✅ Insumo '{nombre}' agregado.")
                    else:
                        st.warning("Escribe el nombre del insumo.")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab3:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT id_insumo, nombre_insumo, stock_actual FROM insumos ORDER BY nombre_insumo")
            insumos = cur.fetchall()
            cur.execute("SELECT id_departamento, nombre FROM departamentos ORDER BY nombre")
            deptos = cur.fetchall()
            cur.close()
            opciones_ins = {f"{i[1]} (stock: {i[2]})": (i[0], i[2]) for i in insumos}
            sel_ins = st.selectbox("Selecciona el insumo", list(opciones_ins.keys()), key="sum_ins_sel")
            id_insumo, stock_actual = opciones_ins[sel_ins]
            depto_sel = st.selectbox("Departamento solicitante", [d[1] for d in deptos], key="sum_depto_sel")
            id_depto  = next(d[0] for d in deptos if d[1] == depto_sel)
            cantidad  = st.number_input("Cantidad a entregar", min_value=1, value=1, key="sum_cant")
            if st.button("📤 Registrar entrega", key="sum_btn"):
                if cantidad > stock_actual:
                    st.error(f"❌ Stock insuficiente. Stock actual: {stock_actual}")
                else:
                    cur2 = get_conn().cursor()
                    cur2.execute("""
                        INSERT INTO suministros (id_insumo, id_departamento, fecha_solicitud,
                                                cantidad_requerida, cantidad_entregada, estatus)
                        VALUES (%s, %s, CURRENT_DATE, %s, %s, %s)
                    """, (id_insumo, id_depto, cantidad, cantidad, "entregado"))
                    cur2.execute("UPDATE insumos SET stock_actual = stock_actual - %s WHERE id_insumo = %s",
                                 (cantidad, id_insumo))
                    get_conn().commit()
                    cur2.close()
                    st.success(f"✅ Entrega registrada. Nuevo stock: {stock_actual - cantidad}")
                    st.rerun()
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

# ════════════════════════════════════════
# DIRECCIONAMIENTO IP
# ════════════════════════════════════════
elif menu == "🌐 Direccionamiento IP":
    st.subheader("🌐 Gestion de Direccionamiento IP - IMJUVE")
    tab1, tab2 = st.tabs(["Rangos y Disponibilidad", "Gestionar IPs"])

    with tab1:
        try:
            cur = get_conn().cursor()
            st.markdown("### Rangos y Disponibilidad")
            cur.execute("""
                SELECT
                    r.area_nombre,
                    r.ip_inicial || ' a ' || r.ip_final,
                    r.capacidad_total,
                    COUNT(i.ip) FILTER (WHERE i.estatus ILIKE 'Ocupada%'),
                    COUNT(i.ip) FILTER (WHERE i.estatus ILIKE 'Libre%'),
                    COUNT(i.ip) FILTER (WHERE i.estatus ILIKE 'En Conflicto%')
                FROM cat_rangos_ips r
                LEFT JOIN inventario_ips_completo i
                    ON UPPER(TRIM(r.area_nombre)) = UPPER(TRIM(i.departamento_pestana))
                GROUP BY r.area_nombre, r.ip_inicial, r.ip_final, r.capacidad_total
                ORDER BY r.area_nombre
            """)
            resumen_data = cur.fetchall()
            cur.close()

            if resumen_data:
                df_r = pd.DataFrame(resumen_data, columns=["Area", "Rango de IPs", "Total IPs", "Ocupadas", "Libres", "En Conflicto"])
                df_r['%_num'] = (df_r['Ocupadas'] / df_r['Total IPs'] * 100).round(1)
                df_r['% Usado'] = df_r['%_num'].astype(str) + '%'

                def color_semaforo(val):
                    try:
                        num = float(val.replace('%', ''))
                        if num >= 90:   color = '#ff4b4b'
                        elif num >= 70: color = '#ffa500'
                        else:           color = '#09ab3b'
                        return f'background-color: {color}; color: white; font-weight: bold'
                    except Exception:
                        return ''

                df_styled = df_r.drop(columns=['%_num']).style.map(color_semaforo, subset=['% Usado'])
                st.table(df_styled)
            else:
                st.info("No hay datos en el catalogo de rangos. Verifica 'cat_rangos_ips'.")
        except Exception as e:
            safe_rollback()
            st.error(f"Error: {e}")

    with tab2:
        try:
            cur = get_conn().cursor()
            cur.execute("SELECT DISTINCT departamento_pestana FROM inventario_ips_completo WHERE departamento_pestana IS NOT NULL ORDER BY departamento_pestana")
            areas_db = [a[0] for a in cur.fetchall()]
            cur.close()

            col_f1, col_f2, col_f3 = st.columns([1, 1, 2])
            area_sel    = col_f1.selectbox("Area", ["Todas"] + areas_db, key="gest_ip_area")
            estatus_sel = col_f2.selectbox("Estatus", ["Todos", "Libre", "Ocupada", "En Conflicto"], key="gest_ip_est")
            busqueda    = col_f3.text_input("Buscar IP, usuario, MAC, marca...", placeholder="Ej: 172.17... o Juan...", key="gest_ip_busq")

            query = """
                SELECT ip, usuario, tipo_equipo, institucional_o_personal,
                       mac, marca, modelo, serie, estatus, departamento_pestana, observaciones
                FROM inventario_ips_completo WHERE 1=1
            """
            params = []
            if area_sel != "Todas":
                query += " AND departamento_pestana = %s"
                params.append(area_sel)
            if estatus_sel != "Todos":
                query += " AND estatus ILIKE %s"
                params.append(f"{estatus_sel}%")
            if busqueda:
                query += " AND (ip ILIKE %s OR usuario ILIKE %s OR mac ILIKE %s OR marca ILIKE %s OR serie ILIKE %s)"
                t = f"%{busqueda}%"
                params.extend([t, t, t, t, t])
            query += " ORDER BY departamento_pestana, ip"

            cur2 = get_conn().cursor()
            cur2.execute(query, params)
            cols = ["ip","usuario","tipo_equipo","institucional_o_personal",
                    "mac","marca","modelo","serie","estatus","departamento_pestana","observaciones"]
            df = pd.DataFrame(cur2.fetchall(), columns=cols)
            cur2.close()

            st.caption(f"{len(df)} IPs encontradas — edita directamente en la tabla y presiona **Guardar cambios**. Para liberar una IP cambia el Estatus a **Libre**.")

            TIPOS = ["PC", "Laptop", "Impresora", "Servidor", "Switch", "Camara", "Otro"]
            USOS  = ["Institucional", "Personal"]
            ESTAT = ["Libre", "Ocupada", "En Conflicto"]

            df_edited = st.data_editor(
                df,
                use_container_width=True,
                hide_index=True,
                num_rows="fixed",
                key="ip_data_editor",
                column_config={
                    "ip":                      st.column_config.TextColumn("IP", disabled=True),
                    "departamento_pestana":    st.column_config.TextColumn("Area", disabled=True),
                    "usuario":                 st.column_config.TextColumn("Usuario"),
                    "tipo_equipo":             st.column_config.SelectboxColumn("Tipo Equipo", options=TIPOS),
                    "institucional_o_personal":st.column_config.SelectboxColumn("Uso", options=USOS),
                    "mac":                     st.column_config.TextColumn("MAC"),
                    "marca":                   st.column_config.TextColumn("Marca"),
                    "modelo":                  st.column_config.TextColumn("Modelo"),
                    "serie":                   st.column_config.TextColumn("Serie"),
                    "estatus":                 st.column_config.SelectboxColumn("Estatus", options=ESTAT, required=True),
                    "observaciones":           st.column_config.TextColumn("Observaciones"),
                },
            )

            col_btn, col_exp, col_full = st.columns([1, 1, 2])
            with col_btn:
                guardar = st.button("💾 Guardar cambios", type="primary", use_container_width=True)
            with col_exp:
                if not df_edited.empty:
                    _df_exp = df_edited.rename(columns={
                        "ip": "IP", "usuario": "Usuario", "tipo_equipo": "Tipo Equipo",
                        "institucional_o_personal": "Uso", "mac": "MAC", "marca": "Marca",
                        "modelo": "Modelo", "serie": "Serie", "estatus": "Estatus",
                        "departamento_pestana": "Area", "observaciones": "Observaciones",
                    })
                    excel_data = generar_excel_formateado(_df_exp, "IPs")
                    st.download_button("📥 Exportar filtro", data=excel_data.getvalue(),
                                       file_name=f"ips_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       use_container_width=True)
            with col_full:
                try:
                    cur_all = get_conn().cursor()
                    cur_all.execute("""
                        SELECT ip, usuario, tipo_equipo, institucional_o_personal,
                               mac, marca, modelo, serie, estatus,
                               departamento_pestana, observaciones
                        FROM inventario_ips_completo
                        ORDER BY departamento_pestana,
                                 CASE WHEN estatus = 'Ocupada'      THEN 1
                                      WHEN estatus = 'Libre'        THEN 2
                                      ELSE 3 END,
                                 ip
                    """)
                    rows_all = cur_all.fetchall()
                    cur_all.close()

                    if rows_all:
                        from openpyxl import Workbook
                        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                        from openpyxl.utils import get_column_letter

                        COLS_DISPLAY = {
                            "ip": "IP", "usuario": "Usuario", "tipo_equipo": "Tipo Equipo",
                            "institucional_o_personal": "Uso", "mac": "MAC",
                            "marca": "Marca", "modelo": "Modelo", "serie": "Serie",
                            "estatus": "Estatus", "departamento_pestana": "Area",
                            "observaciones": "Observaciones",
                        }
                        col_keys = list(COLS_DISPLAY.keys())
                        col_headers = list(COLS_DISPLAY.values())

                        df_all = pd.DataFrame(rows_all, columns=col_keys)
                        areas = df_all["departamento_pestana"].dropna().unique().tolist()

                        AZUL    = "1A3C5E"
                        CELESTE = "EAF1FB"
                        NARANJA = "FFF2CC"
                        BORDE   = "BFBFBF"
                        borde = Border(
                            left=Side(style="thin", color=BORDE),
                            right=Side(style="thin", color=BORDE),
                            top=Side(style="thin", color=BORDE),
                            bottom=Side(style="thin", color=BORDE),
                        )
                        fill_libre   = PatternFill("solid", fgColor=CELESTE)
                        fill_ocup    = PatternFill("solid", fgColor="FFFFFF")
                        fill_header  = PatternFill("solid", fgColor=AZUL)

                        wb = Workbook()
                        wb.remove(wb.active)

                        for area in sorted(areas):
                            df_area = df_all[df_all["departamento_pestana"] == area].copy()
                            nombre_hoja = area[:31]
                            ws = wb.create_sheet(title=nombre_hoja)

                            for ci, h in enumerate(col_headers, 1):
                                c = ws.cell(row=1, column=ci, value=h)
                                c.font      = Font(bold=True, color="FFFFFF", size=10)
                                c.fill      = fill_header
                                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                                c.border    = borde
                            ws.row_dimensions[1].height = 26
                            ws.freeze_panes = "A2"

                            for ri, (_, row) in enumerate(df_area.iterrows(), 2):
                                es_libre = str(row.get("estatus", "")).strip().lower() == "libre"
                                for ci, key in enumerate(col_keys, 1):
                                    val = row[key]
                                    if val is None or (isinstance(val, float) and pd.isna(val)):
                                        val = ""
                                    c = ws.cell(row=ri, column=ci, value=val)
                                    c.font      = Font(size=9)
                                    c.alignment = Alignment(vertical="center")
                                    c.border    = borde
                                    c.fill      = fill_libre if es_libre else fill_ocup

                            for ci, (key, header) in enumerate(COLS_DISPLAY.items(), 1):
                                vals = [header] + [
                                    str(r[key]) if r[key] is not None and not (isinstance(r[key], float) and pd.isna(r[key])) else ""
                                    for _, r in df_area.iterrows()
                                ]
                                ws.column_dimensions[get_column_letter(ci)].width = min(max(len(v) for v in vals) + 4, 40)

                        buf_all = BytesIO()
                        wb.save(buf_all)
                        buf_all.seek(0)

                        st.download_button(
                            "📊 Descargar TODAS las IPs (por area)",
                            data=buf_all.getvalue(),
                            file_name=f"IPs_completo_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                            type="secondary",
                        )
                except Exception as e:
                    st.error(f"Error al generar Excel completo: {e}")

            if guardar:
                def limpiar(v):
                    return None if (v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ('', 'nan', 'None')) else str(v).strip()

                try:
                    cur3 = get_conn().cursor()
                    count = 0
                    for _, row in df_edited.iterrows():
                        estatus = limpiar(row.get('estatus')) or 'Libre'
                        if estatus.lower() == 'libre':
                            cur3.execute("""
                                UPDATE inventario_ips_completo
                                SET usuario=NULL, tipo_equipo=NULL, institucional_o_personal=NULL,
                                    mac=NULL, marca=NULL, modelo=NULL, serie=NULL,
                                    estatus='Libre', observaciones=NULL
                                WHERE ip=%s
                            """, (row['ip'],))
                        else:
                            cur3.execute("""
                                UPDATE inventario_ips_completo
                                SET usuario=%s, tipo_equipo=%s, institucional_o_personal=%s,
                                    mac=%s, marca=%s, modelo=%s, serie=%s,
                                    estatus=%s, observaciones=%s
                                WHERE ip=%s
                            """, (
                                limpiar(row.get('usuario')),
                                limpiar(row.get('tipo_equipo')),
                                limpiar(row.get('institucional_o_personal')),
                                limpiar(row.get('mac')),
                                limpiar(row.get('marca')),
                                limpiar(row.get('modelo')),
                                limpiar(row.get('serie')),
                                estatus,
                                limpiar(row.get('observaciones')),
                                row['ip'],
                            ))
                        count += 1
                    get_conn().commit()
                    cur3.close()
                    st.success(f"✅ {count} IP(s) guardadas correctamente.")
                    st.rerun()
                except Exception as e:
                    safe_rollback()
                    st.error(f"Error al guardar: {e}")

        except Exception as e:
            safe_rollback()
            st.error(f"Error en el sistema: {e}")
            st.warning("Verifica la conexion y las tablas en pgAdmin.")
